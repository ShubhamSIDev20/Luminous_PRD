#!/usr/bin/env python3
"""Generate ME_Primary_SRS_V0.1.xlsx from me_srs_data.py.

Re-runnable: deletes and rewrites the workbook every time, so the data module
is the single source of truth. Never hand-edit the .xlsx.
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import me_srs_data as D  # noqa: E402

OUT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "ME_Primary_SRS_V0.1.xlsx",
)

# ---------------------------------------------------------------- palette
INK = "1F2937"
NAVY = "1B3A5C"
NAVY_D = "12283F"
BAND1 = "DCE6F1"
BAND2 = "EAF1F8"
BAND3 = "F4F8FB"
ROW_A = "FFFFFF"
ROW_B = "F7F9FC"
GREY = "6B7280"
NEW_FILL = "FFF4CE"
MOD_FILL = "E3F2E1"
TBD_FILL = "FDE2E2"
HDR_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
BASE = Font(name="Calibri", size=10, color=INK)
BASE_B = Font(name="Calibri", size=10, bold=True, color=INK)
MONO = Font(name="Consolas", size=9, color=INK)

_hair = Side(style="thin", color="BFCBD9")
BOX = Border(left=_hair, right=_hair, top=_hair, bottom=_hair)

TOP_WRAP = Alignment(vertical="top", horizontal="left", wrap_text=True)
TOP_CTR = Alignment(vertical="top", horizontal="center", wrap_text=True)
MID_CTR = Alignment(vertical="center", horizontal="center", wrap_text=True)
MID_L = Alignment(vertical="center", horizontal="left", wrap_text=True)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def put(ws, r, c, val, font=BASE, align=TOP_WRAP, bg=None, border=BOX):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = font
    cell.alignment = align
    if bg:
        cell.fill = fill(bg)
    if border:
        cell.border = border
    return cell


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def banner(ws, ncols, title, subtitle):
    """Two merged rows at the top of every sheet."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = put(ws, 1, 1, title, Font(name="Calibri", size=14, bold=True, color="FFFFFF"),
            MID_L, NAVY_D, None)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    put(ws, 2, 1, subtitle, Font(name="Calibri", size=9, italic=True, color="FFFFFF"),
        MID_L, NAVY, None)
    ws.row_dimensions[2].height = 18
    return c


def header_row(ws, r, headers, ncols=None):
    for i, h in enumerate(headers, start=1):
        put(ws, r, i, h, HDR_FONT, MID_CTR, NAVY)
    ws.row_dimensions[r].height = 34


def est_height(texts_widths, minimum=15):
    """Approximate the row height Excel needs, since openpyxl cannot autofit."""
    lines = 1
    for text, width in texts_widths:
        if not text:
            continue
        n = 0
        for para in str(text).split("\n"):
            n += max(1, -(-len(para) // max(8, int(width * 1.05))))
        lines = max(lines, n)
    return max(minimum, min(409, lines * 13.2 + 4))


# ================================================================ SHEET 1
def sheet_readme(wb):
    ws = wb.create_sheet("0. Read Me")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 3, "B": 26, "C": 104, "D": 3})
    banner(ws, 4, D.TITLE, D.SUBTITLE)
    r = 4
    for kind, a, b in D.README:
        if kind == "H":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            put(ws, r, 2, a, Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
                MID_L, NAVY, None)
            put(ws, r, 4, None, border=None)
            ws.row_dimensions[r].height = 22
            r += 1
        elif kind == "P":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            put(ws, r, 2, a, BASE, TOP_WRAP, None, None)
            ws.row_dimensions[r].height = est_height([(a, 128)])
            r += 1
        elif kind == "KV":
            put(ws, r, 2, a, BASE_B, TOP_WRAP, BAND3)
            put(ws, r, 3, b, BASE, TOP_WRAP)
            ws.row_dimensions[r].height = est_height([(b, 100)])
            r += 1
        elif kind == "GAP":
            ws.row_dimensions[r].height = 7
            r += 1

    # ---- auto-computed summary, so the numbers on this page cannot go stale
    ws.row_dimensions[r].height = 7
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    put(ws, r, 2, "6.  At a glance  (recomputed on every build)",
        Font(name="Calibri", size=11, bold=True, color="FFFFFF"), MID_L, NAVY, None)
    ws.row_dimensions[r].height = 22
    r += 1
    for label, value in summary_rows():
        put(ws, r, 2, label, BASE_B, TOP_WRAP, BAND3)
        put(ws, r, 3, value, BASE, TOP_WRAP)
        ws.row_dimensions[r].height = 15
        r += 1
    ws.freeze_panes = "A4"
    return ws


def summary_rows():
    reqs = [x for x in D.REQUIREMENTS if x[0] == "R"]
    def n(idx, val):
        return sum(1 for x in reqs if x[idx] == val)
    binding = sum(1 for x in reqs if x[3] == "Requirement")
    tbds = sum(1 for x in reqs if "<TBD-" in str(x[2]))
    hil = n(10, "HIL Test")
    return [
        ("Total rows on sheet 3", str(len(reqs))),
        ("  of which binding requirements", str(binding)),
        ("  of which information / description / constraint",
         str(len(reqs) - binding)),
        ("By owner",
         "BM only: %d      Primary only: %d      Both: %d"
         % (n(5, "BM"), n(5, "Primary"), n(5, "BM + Primary"))),
        ("By phase", "Phase 1: %d      Phase 2 (deferred): %d"
         % (n(6, "Phase 1"), n(6, "Phase 2"))),
        ("By change from BTS",
         "Carried Over: %d      Modified: %d      New: %d"
         % (n(9, "Carried Over"), n(9, "Modified"), n(9, "New"))),
        ("By cardinality",
         "Per System: %d      Per Secondary: %d      Per Channel: %d      "
         "Per CAN Port: %d      Per Program: %d"
         % (n(7, "Per System"), n(7, "Per Secondary"), n(7, "Per Channel"),
            n(7, "Per CAN Port"), n(7, "Per Program"))),
        ("Needs real hardware to verify", "%d requirements marked HIL Test" % hil),
        ("Rows containing an unanswered <TBD>", str(tbds)),
        ("BTS V1.7 rows dispositioned on sheet 4",
         "%d of 282 — none unaccounted for" % len(D.TRACE)),
        ("Open questions on sheet 5",
         "%d total  (%d numeric TBDs, %d open issues)"
         % (len(D.ISSUES), sum(1 for i in D.ISSUES if i[1] == "TBD"),
            sum(1 for i in D.ISSUES if i[1] == "Open Issue"))),
        ("High-priority blockers",
         "%d — start with these"
         % sum(1 for i in D.ISSUES if i[7] == "High")),
    ]


# ================================================================ SHEET 2
def sheet_revision(wb):
    ws = wb.create_sheet("1. Revision History")
    ws.sheet_view.showGridLines = False
    hdr = ["Version", "Date", "Author", "Reviewed By", "Approved By", "Summary of Change"]
    widths(ws, {"A": 11, "B": 13, "C": 20, "D": 20, "E": 20, "F": 86})
    banner(ws, len(hdr), "Revision History", D.SUBTITLE)
    header_row(ws, 4, hdr)
    r = 5
    for row in D.REVISIONS:
        for i, v in enumerate(row, start=1):
            put(ws, r, i, v, BASE, TOP_WRAP if i == 6 else TOP_CTR,
                ROW_A if r % 2 else ROW_B)
        ws.row_dimensions[r].height = est_height([(row[5], 84)])
        r += 1
    ws.freeze_panes = "A5"
    return ws


# ================================================================ SHEET 3
def sheet_topology(wb):
    ws = wb.create_sheet("2. Topology and Terms")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 3, "B": 30, "C": 52, "D": 52, "E": 3})
    banner(ws, 5, "System Topology, Scope Boundary and Terminology", D.SUBTITLE)
    r = 4
    for kind, a, b, c in D.TOPOLOGY:
        if kind == "H":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
            put(ws, r, 2, a, Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
                MID_L, NAVY, None)
            ws.row_dimensions[r].height = 22
            r += 1
        elif kind == "P":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
            put(ws, r, 2, a, BASE, TOP_WRAP, None, None)
            ws.row_dimensions[r].height = est_height([(a, 130)])
            r += 1
        elif kind == "TH":
            for i, v in enumerate((a, b, c), start=2):
                put(ws, r, i, v, HDR_FONT, MID_CTR, NAVY)
            ws.row_dimensions[r].height = 20
            r += 1
        elif kind == "TR":
            put(ws, r, 2, a, BASE_B, TOP_WRAP, BAND3)
            put(ws, r, 3, b, BASE, TOP_WRAP)
            put(ws, r, 4, c, BASE, TOP_WRAP)
            ws.row_dimensions[r].height = est_height([(b, 50), (c, 50)])
            r += 1
        elif kind == "ART":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
            put(ws, r, 2, a, MONO, TOP_WRAP, BAND3, None)
            ws.row_dimensions[r].height = est_height([("x" * 1, 130)],
                                                     minimum=13.2 * (a.count("\n") + 1) + 6)
            r += 1
        elif kind == "GAP":
            ws.row_dimensions[r].height = 7
            r += 1
    ws.freeze_panes = "A4"
    return ws


# ================================================================ SHEET 4
REQ_HDR = [
    ("No", 6), ("Requirement ID", 16), ("Requirement Statement", 92),
    ("Type", 13), ("Category", 15), ("Applicability", 15), ("Phase", 9),
    ("Cardinality", 15), ("Traces To (BTS)", 16), ("Change", 13),
    ("Verification", 13), ("Ref", 18), ("Status", 13), ("Remarks / Rationale", 46),
]

CHANGE_BG = {"New": NEW_FILL, "Modified": MOD_FILL}


def sheet_requirements(wb):
    ws = wb.create_sheet("3. ME SW Requirements")
    ws.sheet_view.showGridLines = False
    ncols = len(REQ_HDR)
    widths(ws, {get_column_letter(i): w for i, (_, w) in enumerate(REQ_HDR, start=1)})
    banner(ws, ncols, "ME Phase 1 — Software Requirements (Primary Board and Battery Manager)",
           D.SUBTITLE)
    header_row(ws, 4, [h for h, _ in REQ_HDR])

    r = 5
    seq = 0
    seen = set()
    dup = []
    for item in D.REQUIREMENTS:
        if item[0] == "S":
            _, level, title = item
            bg = {1: BAND1, 2: BAND2, 3: BAND3}[level]
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            put(ws, r, 1, title,
                Font(name="Calibri", size=10 + (1 if level == 1 else 0), bold=True,
                     color=NAVY_D), MID_L, bg)
            ws.row_dimensions[r].height = 20 if level == 1 else 17
            ws.row_dimensions[r].outlineLevel = 0
            r += 1
            continue

        _, rid, stmt, typ, cat, app, phase, card, traces, change, verif, ref, status, rem = item
        if rid in seen:
            dup.append(rid)
        seen.add(rid)
        seq += 1
        base = ROW_A if seq % 2 else ROW_B
        vals = [seq, rid, stmt, typ, cat, app, phase, card, traces, change,
                verif, ref, status, rem]
        for i, v in enumerate(vals, start=1):
            bg = base
            if i == 10:
                bg = CHANGE_BG.get(change, base)
            if i == 3 and "<TBD-" in str(stmt):
                bg = TBD_FILL
            fnt = BASE
            align = TOP_WRAP
            if i == 1:
                align = TOP_CTR
            elif i == 2:
                fnt, align = MONO, TOP_CTR
            elif i in (4, 5, 6, 7, 8, 10, 11, 13):
                align = TOP_CTR
            elif i == 9:
                fnt, align = MONO, TOP_CTR
            elif i == 12:
                fnt = MONO
            if typ == "Information":
                fnt = Font(name="Calibri", size=10, italic=True, color=GREY) \
                    if i == 3 else fnt
            put(ws, r, i, v, fnt, align, bg)
        ws.row_dimensions[r].height = est_height([(stmt, 90), (rem, 44)])
        r += 1

    last = r - 1
    ws.auto_filter.ref = f"A4:{get_column_letter(ncols)}{last}"
    ws.freeze_panes = "C5"

    dvs = {
        4: D.V_TYPE, 5: D.V_CATEGORY, 6: D.V_APPLICABILITY, 7: D.V_PHASE,
        8: D.V_CARDINALITY, 10: D.V_CHANGE, 11: D.V_VERIFICATION, 13: D.V_STATUS,
    }
    for col, opts in dvs.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"',
                            allow_blank=True, showErrorMessage=True)
        dv.error = "Value must be one of: " + ", ".join(opts)
        dv.errorTitle = "Not an allowed value"
        ws.add_data_validation(dv)
        L = get_column_letter(col)
        dv.add(f"{L}5:{L}{last}")

    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws, seq, dup


# ================================================================ SHEET 5
def sheet_trace(wb):
    ws = wb.create_sheet("4. Traceability BTS to ME")
    ws.sheet_view.showGridLines = False
    hdr = [("BTS Req No", 14), ("BTS Applicability", 16), ("BTS Status", 14),
           ("Disposition", 22), ("ME Requirement ID(s)", 30), ("Reason / Note", 80)]
    widths(ws, {get_column_letter(i): w for i, (_, w) in enumerate(hdr, start=1)})
    banner(ws, len(hdr),
           "Traceability — every requirement of BTS_Primary_SW_Requirement_Analysis V1.7 "
           "and what became of it", D.SUBTITLE)
    header_row(ws, 4, [h for h, _ in hdr])
    disp_bg = {
        "Carried Over": "EAF1F8", "Modified": MOD_FILL,
        "Superseded by New": NEW_FILL, "Dropped — Secondary scope": "F0F0F0",
        "Dropped — not applicable": "F0F0F0", "Moved to Open Issues": TBD_FILL,
        "Deferred to Phase 2": "FFF0E0",
    }
    r = 5
    for row in D.TRACE:
        bts, app, st, disp, meids, note = row
        for i, v in enumerate(row, start=1):
            bg = disp_bg.get(disp, ROW_A) if i == 4 else (ROW_A if r % 2 else ROW_B)
            fnt = MONO if i in (1, 5) else BASE
            align = TOP_CTR if i in (1, 2, 3, 4) else TOP_WRAP
            put(ws, r, i, v, fnt, align, bg)
        ws.row_dimensions[r].height = est_height([(note, 78), (meids, 28)])
        r += 1
    ws.auto_filter.ref = f"A4:F{r - 1}"
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"
    return ws, r - 5


# ================================================================ SHEET 6
def sheet_issues(wb):
    ws = wb.create_sheet("5. Open Issues and TBD")
    ws.sheet_view.showGridLines = False
    hdr = [("ID", 11), ("Kind", 12), ("Area", 20), ("Question / Item to resolve", 80),
           ("Why it blocks work", 62), ("Affects Req IDs", 26), ("Owner", 16),
           ("Priority", 10), ("Status", 12)]
    widths(ws, {get_column_letter(i): w for i, (_, w) in enumerate(hdr, start=1)})
    banner(ws, len(hdr),
           "Open Issues and TBD Register — nothing here was silently assumed", D.SUBTITLE)
    header_row(ws, 4, [h for h, _ in hdr])
    prio_bg = {"High": "FDE2E2", "Medium": "FFF4CE", "Low": "EAF1F8"}
    r = 5
    for row in D.ISSUES:
        for i, v in enumerate(row, start=1):
            bg = prio_bg.get(row[7], ROW_A) if i == 8 else (ROW_A if r % 2 else ROW_B)
            fnt = MONO if i in (1, 6) else BASE
            align = TOP_CTR if i in (1, 2, 7, 8, 9) else TOP_WRAP
            put(ws, r, i, v, fnt, align, bg)
        ws.row_dimensions[r].height = est_height([(row[3], 78), (row[4], 60)])
        r += 1
    ws.auto_filter.ref = f"A4:I{r - 1}"
    ws.freeze_panes = "D5"
    ws.print_title_rows = "4:4"
    return ws, r - 5


# ================================================================ SHEET 7
def sheet_refs(wb):
    ws = wb.create_sheet("6. References and Legend")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 3, "B": 22, "C": 62, "D": 52, "E": 3})
    banner(ws, 5, "Source Documents and Column Legend", D.SUBTITLE)
    r = 4
    for kind, a, b, c in D.REFERENCES:
        if kind == "H":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
            put(ws, r, 2, a, Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
                MID_L, NAVY, None)
            ws.row_dimensions[r].height = 22
            r += 1
        elif kind == "TH":
            for i, v in enumerate((a, b, c), start=2):
                put(ws, r, i, v, HDR_FONT, MID_CTR, NAVY)
            ws.row_dimensions[r].height = 20
            r += 1
        elif kind == "TR":
            put(ws, r, 2, a, BASE_B, TOP_WRAP, BAND3)
            put(ws, r, 3, b, BASE, TOP_WRAP)
            put(ws, r, 4, c, BASE, TOP_WRAP)
            ws.row_dimensions[r].height = est_height([(b, 60), (c, 50)])
            r += 1
        elif kind == "GAP":
            ws.row_dimensions[r].height = 7
            r += 1
    ws.freeze_panes = "A4"
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb)
    sheet_revision(wb)
    sheet_topology(wb)
    _, nreq, dup = sheet_requirements(wb)
    _, ntrace = sheet_trace(wb)
    _, nissue = sheet_issues(wb)
    sheet_refs(wb)
    wb.properties.title = D.TITLE
    wb.properties.creator = "Ador Powertron Ltd — ME Project"
    wb.properties.subject = "ME Phase 1 Software Requirements Specification"
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY
        ws.sheet_view.tabSelected = False
    wb.worksheets[0].sheet_view.tabSelected = True
    wb.active = 0
    if D.TRACE_MISSING:
        print("!! BTS rows neither traced nor listed as dropped:", D.TRACE_MISSING)
        return 1
    wb.save(OUT)

    print(f"written : {OUT}")
    print(f"rows    : {nreq} requirement/information rows")
    print(f"trace   : {ntrace} BTS rows dispositioned")
    print(f"issues  : {nissue} open issues / TBDs")
    if dup:
        print("!! DUPLICATE IDS:", dup)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
